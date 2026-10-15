#!/usr/bin/env python3
"""GDF sanity checker: ZIP-safe, long-format aware, Excel report."""
import argparse, csv, io, json, re, zipfile
from pathlib import Path
from collections import defaultdict, Counter
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

VAR = 'Top Level Field Name'
RESP = 'Response Name'
RID = 'Respondent.Serial'
WEEK = 'Week number'
LEVELS = [
    ('First Level Iterator Name', 'First Level Iterator Label'),
    ('Second Level Iterator Name', 'Second Level Iterator Label'),
    ('Third Level Iterator Name', 'Third Level Iterator Label'),
    ('Fourth Level Iterator Name', 'Fourth Level Iterator Label')
]
BRAND_VARS = {
    'BHCAsked', 'BE6Total', 'BE7Total', 'BP1', 'BP6', 'COM19Flip', 'UsageTab',
    'Q7aYes', 'QSLP', 'SLPAT', 'Pack1Core', 'UPOS2_HeardOf', 'upos1', 'upos2'
}
GDF = [
    ('Total Awareness', 'Brand Funnel', 'UsageTab', 'TotalAwareness'),
    ('Consideration - T2B', 'Brand Funnel', 'BE7Total', ''),
    ('Use Nowadays', 'Brand Funnel', 'UsageTab', 'UseNow'),
    ('Use Most Often', 'Brand Funnel', 'UsageTab', 'BuyMostOften'),
    ('Trial', 'Brand Funnel', 'UsageTab', 'EverUsed'),
    ('Bought Last', 'Brand Funnel', 'UsageTab', 'BoughtLast'),
    ('Brand Imagery Statements', 'Brand Imagery', 'upos2', ''),
    ('All available channels', 'Media Awareness', 'COM19Flip', ''),
    ('DBCA', 'Media Awareness', 'COM19Flip', 'Target_Digital_NET'),
    ('TBCA', 'Media Awareness', 'Q7aYes', ''),
    ('Affinity - T2B', 'Underlying Metrics', 'BE6Total', ''),
    ('Meets Needs - T2B', 'Underlying Metrics', 'BP1', 'Needs'),
    ('Dynamic - T2B', 'Underlying Metrics', 'BP1', 'FollowVsTrends'),
    ('Unique - T2B', 'Underlying Metrics', 'BP1', 'SameVsDifferent'),
    ('Weight Variable', 'Weighting', 'WEIGHT', ''),
    ('Attractive Pack', 'Packaging', 'upos2', ''),
    ('Packs Help to Know What to Expect', 'Packaging', 'Pack1Core', 'WhatToExpectFromProduct'),
    ('Packs Stand Out on Shelf', 'Packaging', 'Pack1Core', 'StandOutOnShelf'),
    ('Packs Very Easy to Use', 'Packaging', 'Pack1Core', 'EasyToUse'),
    ('Used Daily', 'Digital Platforms Usage', 'SOCMED2_CORE', ''),
    ('Used Regularly', 'Digital Platforms Usage', 'SOCMED1_CORE', '')
]
EXT = [
    ('Brand Power Share', 'Brand Equity', 'BEAST'),
    ('Premium Index', 'Brand Equity/Pricing', 'BEAST'),
    ('Salient Index', 'Brand Equity', 'BEAST'),
    ('Meaningful Index', 'Brand Equity', 'BEAST'),
    ('Different Index', 'Brand Equity', 'BEAST'),
    ('Perceived Price Index', 'Pricing', '6P'),
    ('Perceived Worth Index', 'Pricing', '6P'),
    ('Attractive Pack (Total)', 'Additional UBS Metrics', '6P'),
    ('Better Quality Index', 'Additional UBS Metrics', '6P'),
    ('SLP Absolute', 'Sustainability', '6P'),
    ('SLP Index', 'Sustainability', '6P'),
    ('SLP Profile', 'Sustainability', '6P'),
    ('Behave Socially Responsible', 'Sustainability', '6P'),
    ('Environmentally Responsible', 'Sustainability', '6P'),
    ('Packaging Better for Environment', 'Sustainability', '6P'),
    ('Recycled Packaging', 'Sustainability', '6P'),
    ('Recycled Packaging Index', 'Sustainability', '6P')
]

def detect(raw):
    e = 'utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else ('utf-8-sig' if raw.startswith(b'\xef\xbb\xbf') else 'utf-8')
    t = raw[:65536].decode(e, errors='replace')
    try:
        s = csv.Sniffer().sniff(t, delimiters=',;\t|').delimiter
    except:
        s = ','
    return e, s

def files(paths):
    for x in paths:
        p = Path(x)
        for q in (p.rglob('*') if p.is_dir() else [p]):
            if not q.is_file(): continue
            if q.suffix.lower() == '.csv':
                yield q.name, q.read_bytes(), q.stat().st_size
            elif q.suffix.lower() == '.zip':
                with zipfile.ZipFile(q) as z:
                    for i in z.infolist():
                        if not i.is_dir() and i.filename.lower().endswith('.csv'):
                            yield q.name + '::' + i.filename, z.read(i.filename), i.file_size

def cellname(n):
    m = re.match(r'UB([A-Z]{2})([A-Z0-9]{4})_', Path(n.split('::')[-1]).stem, re.I)
    return (m.group(1).upper(), m.group(2).upper()) if m else ('', '')

def norm(x):
    return re.sub('[^a-z0-9]', '', x.lower())

def main():
    a = argparse.ArgumentParser()
    a.add_argument('inputs', nargs='+')
    a.add_argument('--filter-workbook', required=True)
    a.add_argument('--sheet', default='Scott list of needed filters')
    a.add_argument('-o', '--output', default='GDF_Sanity_Report.xlsx')
    z = a.parse_args()
    
    # Read Excel layout safely
    sf = pd.read_excel(z.filter_workbook, sheet_name=z.sheet, dtype=str, engine='openpyxl').fillna('')
    varcol = 'Variable if already available in GDF files'
    
    actual_col = None
    for col in sf.columns:
        if str(col).strip().lower() == varcol.lower():
            actual_col = col
            break
            
    if actual_col is None:
        raise KeyError(f"Could not find target column header. Expected something close to: '{varcol}'. Found columns: {list(sf.columns)}")
        
    # Safeguard string serialization values cleanly
    expected = sorted(set(str(item).strip() for item in sf[actual_col] if str(item).strip() not in ['', 'nan']))
    
    fr, cells, issues, dups, vp, metrics, occ = [], [], [], [], [], [], []
    code_names = defaultdict(set)
    name_codes = defaultdict(set)
    source = list(files(z.inputs))
    med = pd.Series([s for n, r, s in source if 'CASEDATA' in n.upper()]).median()
    
    for fn, raw, size in source:
        typ = next((x for x in ['CASEDATA', 'METADATA', 'EXISTENCE'] if x in fn.upper()), 'OTHER')
        e, sep = detect(raw)
        c0, k0 = cellname(fn)
        rows = 0
        try:
            head = pd.read_csv(io.BytesIO(raw), encoding=e, sep=sep, nrows=5, dtype=str, keep_default_na=False, on_bad_lines='warn')
            cols = len(head.columns)
        except Exception as ex:
            fr.append([fn, c0, k0, typ, size / 1024, 0, 0, e, sep, 'FAIL'])
            continue
            
        if typ != 'CASEDATA':
            for ch in pd.read_csv(io.BytesIO(raw), encoding=e, sep=sep, chunksize=100000, dtype=str, keep_default_na=False, on_bad_lines='warn'):
                rows += len(ch)
            fr.append([fn, c0, k0, typ, round(size / 1024, 1), rows, cols, e, sep, 'PASS'])
            continue
            
        vars = set()
        resp = defaultdict(set)
        cs, ks, waves, ids, rw = set(), set(), set(), set(), set()
        dup = 0
        wc = Counter()
        wv = defaultdict(set)
        brands = defaultdict(lambda: defaultdict(set))
        
        for ch in pd.read_csv(io.BytesIO(raw), encoding=e, sep=sep, chunksize=100000, dtype=str, keep_default_na=False, on_bad_lines='warn'):
            ch = ch.fillna('')
            rows += len(ch)
            dup += int(ch.duplicated().sum())
            vars.update(ch[VAR].unique())
            
            for v, g in ch.groupby(VAR):
                resp[v].update(g[RESP].unique())
                
            cs.update(ch.loc[ch[VAR].eq('Country'), RESP])
            ks.update(ch.loc[ch[VAR].eq('Category'), RESP])
            waves.update(ch[WEEK])
            ids.update(ch[RID])
            rw.update(map(tuple, ch[[RID, WEEK]].drop_duplicates().values))
            
            for _, r in ch[ch[VAR].isin(['WEIGHT', 'weightage', 'Weight_human', 'boost_factor', 'ESSBase', 'human_weights_squared'])].iterrows():
                q = (r[RID], r[WEEK], r[VAR])
                wc[q] += 1
                wv[q].add(r[RESP])
                
            for _, r in ch[ch.duplicated(keep=False)].head(max(0, 25 - len(dups))).iterrows():
                dups.append([fn, c0, k0, r[RID], r[WEEK], r[VAR], r[RESP]])
                
            b = ch[ch[VAR].isin(BRAND_VARS)]
            for cc, ll in LEVELS:
                for _, r in b[[cc, ll, WEEK]].drop_duplicates().iterrows():
                    code, label, w = r[cc], r[ll], r[WEEK]
                    if re.fullmatch(r'_\d+[A-Za-z]?', code) and label:
                        brands[code][label].add(w)
                        code_names[code].add(label)
                        name_codes[norm(label)].add(code)
                        occ.append([fn, c0, k0, code, label, w])
                        
        cs.discard('')
        ks.discard('')
        c = '|'.join(sorted(cs)) or c0
        k = '|'.join(sorted(ks)) or k0
        
        missing = [v for v in expected if str(v).split(' - ')[0].strip() not in vars]
        mm = []
        for dash, grp, fld, detail in GDF:
            ok = fld in vars
            metrics.append([fn, c, k, 'GDF metric', dash, grp, 'GDF', fld, detail, 'Yes' if ok else 'No', 'PASS' if ok else 'FAIL', '' if ok else 'Fixed field code missing'])
            if not ok: mm.append(fld)
            
        for v in expected:
            prefix = str(v).split(' - ')[0].strip()
            vp.append([fn, c, k, v, 'Yes' if prefix in vars else 'No', 'PASS' if prefix in vars else 'FAIL'])
            
        dupw = sum(n > 1 for q, n in wc.items() if q[2] == 'WEIGHT')
        conf = sum(len(v) > 1 for q, v in wv.items() if q[2] == 'WEIGHT')
        miss = sum((r, w, 'WEIGHT') not in wc for r, w in rw)
        bad = 0
        for q, vals in wv.items():
            if q[2] == 'WEIGHT':
                for v in vals:
                    try: bad += float(v) <= 0
                    except: bad += 1
                    
        st = 'FAIL' if missing or mm or dup or dupw or conf or miss or bad or len(cs) != 1 or len(ks) != 1 else 'PASS'
        cells.append([fn, c, k, ', '.join(sorted(waves)), rows, len(ids), len(vars), len(brands), len(missing), len(set(mm)), dup, dupw + conf + miss + bad, st])
        issues += [
            [fn, c, k, 'Cell', 'Single country and category', 'PASS' if len(cs) == 1 and len(ks) == 1 else 'FAIL', 'High', f'Countries={sorted(cs)}; Categories={sorted(ks)}'],
            [fn, c, k, 'Respondent', 'Exact duplicate rows', 'PASS' if not dup else 'FAIL', 'High', dup],
            [fn, c, k, 'Weight', 'Completeness and uniqueness', 'PASS' if not (dupw + conf + miss + bad) else 'FAIL', 'High', f'duplicate={dupw}; conflict={conf}; missing={miss}; invalid={bad}'],
            [fn, c, k, 'Variables', 'Scott filters', 'PASS' if not missing else 'FAIL', 'Medium', 'Missing: ' + str(missing)],
            [fn, c, k, 'Variables', 'Fixed GDF fields', 'PASS' if not mm else 'FAIL', 'High', 'Missing: ' + str(sorted(set(mm)))]
        ]
        fr.append([fn, c, k, typ, round(size / 1024, 1), rows, cols, e, 'TAB' if sep == '\t' else sep, 'WARN' if med and size < .2 * med else 'PASS'])
        
    for d, g, s in EXT:
        metrics.append(['Not evaluated in GDF audit', '', '', 'External metric reference', d, g, s, '', '', 'Not checked', 'INFO', f'Supplied by {s}'])
        
    od = pd.DataFrame(occ, columns=['File', 'Country', 'Category', 'Brand_Code', 'Brand_Name', 'Wave']).drop_duplicates()
    br = []
    if len(od):
        for (c, k, code), g in od.groupby(['Country', 'Category', 'Brand_Code']):
            names = sorted(g.Brand_Name.unique())
            ww = sorted(g.Wave.unique())
            mis = len(code_names[code]) > 1 or any(len(name_codes[norm(n)]) > 1 for n in names)
            com = 'Review as harmonization issue: ' if mis else ''
            com += ('same code has names ' + str(sorted(code_names[code]))) if len(code_names[code]) > 1 else ''
            br.append([c, k, code, ' | '.join(names), ww[0], ww[-1], '', 'Yes' if mis else 'No', 'FAIL' if mis else 'PASS', com])
        bd = pd.DataFrame(br, columns=['Country', 'Category', 'Brand_Code', 'Brand_Name', 'First_Wave_Observed', 'Latest_Wave_Observed', 'New_Brand_After_Cell_First_Wave', 'Name_Mismatch', 'Status', 'Comments'])
        for (c, k), idx in bd.groupby(['Country', 'Category']).groups.items():
            first = min(bd.loc[list(idx), 'First_Wave_Observed'])
            bd.loc[list(idx), 'New_Brand_After_Cell_First_Wave'] = bd.loc[list(idx), 'First_Wave_Observed'].map(lambda x: 'Yes' if x > first else 'No')
    else:
        bd = pd.DataFrame(columns=['Country', 'Category', 'Brand_Code', 'Brand_Name', 'First_Wave_Observed', 'Latest_Wave_Observed', 'New_Brand_After_Cell_First_Wave', 'Name_Mismatch', 'Status', 'Comments'])
        
    sf2 = pd.DataFrame({
        'Scott_Level_1': sf.iloc[:, 0],
        'Filter_Level_1': sf.iloc[:, 1],
        'Filter_Level_2': sf.iloc[:, 2],
        'Source_Status': sf.get('DP to add/Available', ''),
        'Expected_Variable': sf[actual_col]
    })
    sf2['Check_Scope'] = sf2.apply(lambda r: 'Current GDF' if r.Expected_Variable else ('Future DP addition' if 'DP to add' in str(r.Source_Status) else 'Dashboard/Other'), axis=1)
    
    dfs = {
        '01_Summary': pd.DataFrame(cells, columns=['File', 'Country', 'Category', 'Waves', 'Rows', 'Respondents', 'Variables', 'Brands', 'Missing_Demographic_Filters', 'Missing_GDF_Metric_Fields', 'Exact_Duplicate_Rows', 'Weight_Issues', 'Status']),
        '02_File_Checks': pd.DataFrame(fr, columns=['File', 'Country', 'Category', 'Type', 'Size_KB', 'Rows', 'Columns', 'Encoding', 'Delimiter', 'Status']),
        '03_Variable_Checks': sf2,
        '04_Filter_Matrix': pd.DataFrame(vp, columns=['File', 'Country', 'Category', 'Expected_Variable', 'Present', 'Status']),
        '05_Metric_Checks': pd.DataFrame(metrics, columns=['File', 'Country', 'Category', 'Check_Type', 'Dashboard_Variable', 'Metric_Group', 'Data_Source', 'Expected_Field_Code', 'Expected_Dataset_Metric', 'Present', 'Status', 'Comments']),
        '06_Brand_Codes': bd,
        '07_Issues': pd.DataFrame(issues, columns=['File', 'Country', 'Category', 'Area', 'Check', 'Status', 'Severity', 'Details']),
        '08_Duplicate_Examples': pd.DataFrame(dups, columns=['File', 'Country', 'Category', 'Respondent', 'Wave', 'Variable', 'Response']),
        '09_Config': pd.DataFrame([['Brand newness', 'First observed wave later than earliest brand wave in country x category'], ['Brand mismatch', 'One code/multiple names or normalized name/multiple codes'], ['Duplicate key', 'Exact full long-format row'], ['Weight key', 'Respondent.Serial + Week number + WEIGHT'], ['External metrics', '6P and BEAST are reference-only']], columns=['Setting', 'Value'])
    }
    
    with pd.ExcelWriter(z.output, engine='openpyxl') as w:
        for n, d in dfs.items(): d.to_excel(w, sheet_name=n, index=False)
        
    wb = load_workbook(z.output)
    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'
        ws.sheet_view.showGridLines = True
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.fill = PatternFill('solid', fgColor='1F4E78')
            c.font = Font(color='FFFFFF', bold=True)
            c.alignment = Alignment(horizontal='center')
        for j in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(j)].width = min(55, max(11, max([len(str(ws.cell(i, j).value or '')) for i in range(1, min(200, ws.max_row) + 1)] + [8]) + 2))
    wb.save(z.output)
    print(z.output)

if __name__ == '__main__':
    main()
